"""Master benchmark table: 4 observers x 7 conditions (measured & hidden RMSE)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

conds = ["Clean","Moderate noise","Strong noise","Mass worst (±20%)",
         "Inertia worst (+20%)","Arm worst (-20%)","Combined-worst"]
# (measured, hidden) per observer, per condition (same order as conds)
DATA = {
 "PINN":       [(0.0442,0.0559),(0.0775,0.0699),(0.0850,0.0720),(0.1509,0.3461),(0.1005,0.3069),(0.1239,0.3835),(0.3384,0.6963)],
 "Luenberger": [(0.0100,0.0094),(0.0127,0.0213),(0.0183,0.0368),(0.0146,0.1220),(0.0105,0.0355),(0.0109,0.0477),(0.0189,0.2457)],
 "EKF":        [(0.0000,0.0080),(0.0071,0.0131),(0.0110,0.0174),(0.0004,0.4873),(0.0000,0.0485),(0.0001,0.0661),(0.0006,0.6769)],
 "UKF":        [(0.0000,0.0129),(0.0073,0.0163),(0.0113,0.0198),(0.0004,0.4920),(0.0001,0.0557),(0.0001,0.0737),(0.0006,0.6767)],
}
obs = ["PINN","Luenberger","EKF","UKF"]

WHITE=Font(bold=True,color="FFFFFF"); BOLD=Font(bold=True)
PF=PatternFill("solid",fgColor="21295C"); MF=PatternFill("solid",fgColor="1D5FB0")
BEST=PatternFill("solid",fgColor="D6EFD6"); BFONT=Font(bold=True,color="1B5E20")
thin=Side(style="thin",color="C9D2D9"); border=Border(thin,thin,thin,thin); ctr=Alignment(horizontal="center",vertical="center")

wb=openpyxl.Workbook(); ws=wb.active; ws.title="Benchmark"
ws["A1"]="Observer benchmark on the spiral (10 unseen flights) - RMSE"; ws["A1"].font=Font(bold=True,size=14)
ws["A2"]="Classical observers use the live measurement stream; the PINN uses only time + initial condition. Green = best (lowest) per row."
ws["A2"].font=Font(italic=True,color="555555")

ws.merge_cells("A4:A5"); ws["A4"]="Condition"; ws["A4"].font=WHITE; ws["A4"].fill=PF; ws["A4"].alignment=ctr
c=2
for o in obs:
    ws.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+1)
    cell=ws.cell(row=4,column=c,value=o); cell.font=WHITE; cell.fill=PF; cell.alignment=ctr
    for j,lab in enumerate(["meas","hid"]):
        cc=ws.cell(row=5,column=c+j,value=lab); cc.font=WHITE; cc.fill=MF; cc.alignment=ctr; cc.border=border
    c+=2
for rr in range(4,6):
    for cc in range(1,10): ws.cell(row=rr,column=cc).border=border

for i,cond in enumerate(conds):
    r=6+i; a=ws.cell(row=r,column=1,value=cond); a.font=BOLD; a.border=border
    meas_cells=[]; hid_cells=[]
    c=2
    for o in obs:
        m,h=DATA[o][i]
        cm=ws.cell(row=r,column=c,value=m); cm.number_format="0.0000"; cm.alignment=ctr; cm.border=border
        ch=ws.cell(row=r,column=c+1,value=h); ch.number_format="0.0000"; ch.alignment=ctr; ch.border=border
        meas_cells.append(cm); hid_cells.append(ch); c+=2
    mmin=min(meas_cells,key=lambda x:x.value); hmin=min(hid_cells,key=lambda x:x.value)
    mmin.fill=BEST; mmin.font=BFONT; hmin.fill=BEST; hmin.font=BFONT

ws.cell(row=14,column=1,value="Takeaways:").font=BOLD
ws.cell(row=15,column=1,value="- Measured states: EKF/UKF ~0, Luenberger ~0.01 (they use live sensors); PINN far behind.")
ws.cell(row=16,column=1,value="- Hidden under mismatch: EKF/UKF become brittle (~0.68, comparable to PINN 0.70); Luenberger best (0.25).")
ws.cell(row=17,column=1,value="- PINN needs no live measurements at inference - a structural advantage, not raw-accuracy on clean data.")
for k in range(15,18): ws.cell(row=k,column=1).font=Font(italic=True,color="666666")

ws.column_dimensions["A"].width=22
for col in "BCDEFGHI": ws.column_dimensions[col].width=9
ws.freeze_panes="B6"
wb.save("docs/benchmark_master.xlsx")
print("saved -> docs/benchmark_master.xlsx")
