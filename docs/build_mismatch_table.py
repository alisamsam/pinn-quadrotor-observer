"""Rearranged mismatch table: sections Mass/Inertia/Arm, each split V1/V2, rows = deviation %."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

devs = ["-20%", "-10%", "0%", "+10%", "+20%"]
old = {
 "mass":    {"meas":[1.0762,0.5137,0.0286,0.4486,0.8493], "hid":[1.2747,0.5735,0.0306,0.4416,0.7783]},
 "inertia": {"meas":[0.3404,0.1946,0.0286,0.3334,1.2356], "hid":[0.3582,0.2095,0.0306,0.3692,1.3349]},
 "arm":     {"meas":[1.4343,0.4080,0.0286,0.1796,0.2952], "hid":[1.5571,0.4454,0.0306,0.1936,0.3132]},
}
v2 = {
 "mass":    {"meas":[0.1509,0.0892,0.0436,0.0679,0.1068], "hid":[0.3461,0.1758,0.0552,0.0840,0.1119]},
 "inertia": {"meas":[0.0546,0.0462,0.0436,0.0615,0.1005], "hid":[0.1125,0.0823,0.0552,0.1424,0.3069]},
 "arm":     {"meas":[0.1239,0.0650,0.0436,0.0454,0.0519], "hid":[0.3835,0.1598,0.0552,0.0783,0.1034]},
}

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Mismatch table"
WHITE = Font(bold=True, color="FFFFFF"); BOLD = Font(bold=True)
PARAMF = PatternFill("solid", fgColor="21295C")
V1F = PatternFill("solid", fgColor="B23A48"); V2F = PatternFill("solid", fgColor="1D5FB0")
V1L = PatternFill("solid", fgColor="F6D3D3"); V2L = PatternFill("solid", fgColor="D6E6F7")
thin = Side(style="thin", color="AAB4BC"); border = Border(thin,thin,thin,thin)
center = Alignment(horizontal="center", vertical="center")

ws["A1"] = "Model-parameter mismatch — RMSE by deviation (V1 vs V2)"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Frozen nominal-trained observer on flights with perturbed true parameters. meas = measured states, hid = hidden states."
ws["A2"].font = Font(italic=True, color="555555")

params = ["mass","inertia","arm"]
ws.merge_cells("A4:A6"); ws["A4"] = "Deviation"; ws["A4"].font = WHITE
ws["A4"].fill = PARAMF; ws["A4"].alignment = center
col = 2
for p in params:
    c0 = col
    ws.merge_cells(start_row=4,start_column=c0,end_row=4,end_column=c0+3)
    cell = ws.cell(row=4,column=c0,value=p.upper()); cell.font = WHITE; cell.fill = PARAMF; cell.alignment = center
    for name,off,fill in [("Controller V1",0,V1F),("Controller V2",2,V2F)]:
        ws.merge_cells(start_row=5,start_column=c0+off,end_row=5,end_column=c0+off+1)
        cc = ws.cell(row=5,column=c0+off,value=name); cc.font = WHITE; cc.fill = fill; cc.alignment = center
    for j,lab in enumerate(["meas","hid","meas","hid"]):
        cc = ws.cell(row=6,column=c0+j,value=lab); cc.font = BOLD; cc.alignment = center
        cc.fill = V1L if j < 2 else V2L
    col += 4
for r in range(4,7):
    for cc in range(1,14):
        ws.cell(row=r,column=cc).border = border


# data rows
for i, dev in enumerate(devs):
    r = 7 + i
    a = ws.cell(row=r, column=1, value=dev); a.font = BOLD; a.alignment = center; a.border = border
    c0 = 2
    for p in params:
        vals = [old[p]["meas"][i], old[p]["hid"][i], v2[p]["meas"][i], v2[p]["hid"][i]]
        for j, val in enumerate(vals):
            cc = ws.cell(row=r, column=c0+j, value=val)
            cc.number_format = "0.0000"; cc.alignment = center; cc.border = border
            cc.fill = V1L if j < 2 else V2L
        c0 += 4
    if dev == "0%":
        for cc in range(1, 14):
            ws.cell(row=r, column=cc).font = BOLD

ws.cell(row=13, column=1,
        value="Read-down each column: V1 (old) error grows steeply with deviation; V2 stays low = far more robust. 0% = nominal (no mismatch).")
ws.cell(row=13, column=1).font = Font(italic=True, color="666666")

ws.column_dimensions["A"].width = 11
for c in "BCDEFGHIJKLM":
    ws.column_dimensions[c].width = 9
ws.freeze_panes = "B7"
wb.save("docs/mismatch_table_v1_v2.xlsx")
print("saved -> docs/mismatch_table_v1_v2.xlsx")
