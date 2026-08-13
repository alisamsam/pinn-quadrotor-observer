"""Consolidated mismatch comparison: PINN vs Luenberger vs EKF, all 5 deviations."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

devs = ["-20%","-10%","0%","+10%","+20%"]
PINN = {
 "meas": {"mass":[0.1509,0.0892,0.0436,0.0679,0.1068],"inertia":[0.0546,0.0462,0.0436,0.0615,0.1005],"arm":[0.1239,0.0650,0.0436,0.0454,0.0519]},
 "hid":  {"mass":[0.3461,0.1758,0.0552,0.0840,0.1119],"inertia":[0.1125,0.0823,0.0552,0.1424,0.3069],"arm":[0.3835,0.1598,0.0552,0.0783,0.1034]}}
LUEN = {
 "meas": {"mass":[0.0146,0.0121,0.0100,0.0083,0.0098],"inertia":[0.0106,0.0103,0.0100,0.0100,0.0105],"arm":[0.0109,0.0100,0.0100,0.0103,0.0105]},
 "hid":  {"mass":[0.1220,0.0632,0.0094,0.0506,0.1029],"inertia":[0.0254,0.0171,0.0094,0.0159,0.0355],"arm":[0.0477,0.0176,0.0094,0.0164,0.0227]}}
EKF = {
 "meas": {"mass":[0.0004,0.0002,0.0000,0.0002,0.0004],"inertia":[0.0000,0.0000,0.0000,0.0000,0.0000],"arm":[0.0001,0.0000,0.0000,0.0000,0.0000]},
 "hid":  {"mass":[0.4873,0.2456,0.0081,0.2450,0.4843],"inertia":[0.0288,0.0168,0.0081,0.0205,0.0485],"arm":[0.0661,0.0231,0.0081,0.0157,0.0247]}}
COMB = {"PINN":(0.3384,0.6963),"Luenberger":(0.0189,0.2457),"EKF":(0.0006,0.6769)}
methods = [("PINN",PINN),("Luen",LUEN),("EKF",EKF)]
params = ["mass","inertia","arm"]

WHITE=Font(bold=True,color="FFFFFF"); BOLD=Font(bold=True)
PF=PatternFill("solid",fgColor="21295C"); MF=PatternFill("solid",fgColor="1D5FB0")
WORST=PatternFill("solid",fgColor="FFD24D"); WFONT=Font(bold=True,color="8A1C1C")
thin=Side(style="thin",color="C9D2D9"); border=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center")
wb=openpyxl.Workbook(); wb.remove(wb.active)

def sheet(metric,title):
    ws=wb.create_sheet(title)
    ws["A1"]=f"Mismatch — {title} RMSE:  PINN vs Luenberger vs EKF"; ws["A1"].font=Font(bold=True,size=13)
    ws["A2"]="Frozen nominal-trained observer; true drone perturbed. Amber = worst deviation per column."; ws["A2"].font=Font(italic=True,color="555555")
    ws.merge_cells("A4:A5"); ws["A4"]="Deviation"; ws["A4"].font=WHITE; ws["A4"].fill=PF; ws["A4"].alignment=ctr
    c=2
    for p in params:
        ws.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+2)
        cell=ws.cell(row=4,column=c,value=p.upper()); cell.font=WHITE; cell.fill=PF; cell.alignment=ctr
        for j,(mname,_) in enumerate(methods):
            cc=ws.cell(row=5,column=c+j,value=mname); cc.font=WHITE; cc.fill=MF; cc.alignment=ctr; cc.border=border
        c+=3
    for r in range(4,6):
        for cc in range(1,11): ws.cell(row=r,column=cc).border=border
    for i,dv in enumerate(devs):
        r=6+i; a=ws.cell(row=r,column=1,value=dv); a.font=BOLD; a.alignment=ctr; a.border=border
        c=2
        for p in params:
            for j,(_,md) in enumerate(methods):
                cell=ws.cell(row=r,column=c+j,value=md[metric][p][i]); cell.number_format="0.0000"; cell.alignment=ctr; cell.border=border
            c+=3
    # highlight worst per column
    for col in range(2,11):
        rr=max(range(6,11),key=lambda r: ws.cell(row=r,column=col).value)
        ws.cell(row=rr,column=col).fill=WORST; ws.cell(row=rr,column=col).font=WFONT
    ws.column_dimensions["A"].width=11
    for col in "BCDEFGHIJ": ws.column_dimensions[col].width=9
    ws.freeze_panes="B6"

sheet("meas","Measured"); sheet("hid","Hidden")


# combined-worst sheet
ws=wb.create_sheet("Combined-worst")
ws["A1"]="Combined-worst mismatch (mass -20%, inertia +20%, arm -20%)"; ws["A1"].font=Font(bold=True,size=13)
ws["A2"]="Each parameter at its individually-worst direction."; ws["A2"].font=Font(italic=True,color="555555")
for j,h in enumerate(["Method","Measured RMSE","Hidden RMSE"],start=1):
    cell=ws.cell(row=4,column=j,value=h); cell.font=WHITE; cell.fill=MF; cell.alignment=ctr; cell.border=border
for i,(m,(mm,hh)) in enumerate(COMB.items()):
    r=5+i
    ws.cell(row=r,column=1,value=m).font=BOLD
    ws.cell(row=r,column=2,value=mm).number_format="0.0000"
    ws.cell(row=r,column=3,value=hh).number_format="0.0000"
    for c in range(1,4): ws.cell(row=r,column=c).border=border; ws.cell(row=r,column=c).alignment=ctr
ws.cell(row=9,column=1,value="Note: classical observers win on measured (they use live sensors); on hidden under mismatch the EKF becomes brittle (~PINN), Luenberger best.").font=Font(italic=True,color="666666")
for col,w in {"A":14,"B":16,"C":16}.items(): ws.column_dimensions[col].width=w

wb.save("docs/mismatch_all_observers.xlsx")
print("saved -> docs/mismatch_all_observers.xlsx")
