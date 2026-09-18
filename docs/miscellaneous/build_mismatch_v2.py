"""Mismatch comparison: previous controller vs controller_v2. CSV + chart + Excel."""
import csv
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

devs = [-20, -10, 0, 10, 20]
v2 = {
 "mass":    {"meas":[0.1509,0.0892,0.0436,0.0679,0.1068], "hid":[0.3461,0.1758,0.0552,0.0840,0.1119]},
 "inertia": {"meas":[0.0546,0.0462,0.0436,0.0615,0.1005], "hid":[0.1125,0.0823,0.0552,0.1424,0.3069]},
 "arm":     {"meas":[0.1239,0.0650,0.0436,0.0454,0.0519], "hid":[0.3835,0.1598,0.0552,0.0783,0.1034]},
}
v2c = {"meas":0.3384, "hid":0.6963}
old = {
 "mass":    {"meas":[1.0762,0.5137,0.0286,0.4486,0.8493], "hid":[1.2747,0.5735,0.0306,0.4416,0.7783]},
 "inertia": {"meas":[0.3404,0.1946,0.0286,0.3334,1.2356], "hid":[0.3582,0.2095,0.0306,0.3692,1.3349]},
 "arm":     {"meas":[1.4343,0.4080,0.0286,0.1796,0.2952], "hid":[1.5571,0.4454,0.0306,0.1936,0.3132]},
}
oldc = {"meas":3.0034, "hid":3.2302}

# ---- corrected CSV (v2) ----
cols = ["-20%","-10%","nominal","+10%","+20%"]
with open("docs/mismatch_v2_summary.csv","w",newline="") as fp:
    w = csv.writer(fp); w.writerow(["parameter","metric"]+cols)
    for p in ["mass","inertia","arm"]:
        w.writerow([p,"measured RMSE"]+[f"{x:.4f}" for x in v2[p]["meas"]])
        w.writerow([p,"hidden RMSE"]+[f"{x:.4f}" for x in v2[p]["hid"]])
    w.writerow(["combined worst","measured RMSE","","","0.0436","",f"{v2c['meas']:.4f}"])
    w.writerow(["combined worst","hidden RMSE","","","0.0552","",f"{v2c['hid']:.4f}"])
print("saved -> docs/mismatch_v2_summary.csv")

# ---- chart: 2x2 (mass, inertia, arm lines; combined bar) ----
fig, ax = plt.subplots(2,2, figsize=(13,9))
for a,p in zip([ax[0,0],ax[0,1],ax[1,0]], ["mass","inertia","arm"]):
    a.plot(devs, old[p]["hid"], "r-o", lw=2, ms=4, label="old: hidden")
    a.plot(devs, old[p]["meas"],"r--s", lw=1.4, ms=3, label="old: measured")
    a.plot(devs, v2[p]["hid"],  "b-o", lw=2, ms=4, label="v2: hidden")
    a.plot(devs, v2[p]["meas"], "b--s", lw=1.4, ms=3, label="v2: measured")
    a.set_title(f"{p} mismatch"); a.set_xlabel("parameter deviation (%)")
    a.set_ylabel("test RMSE"); a.grid(alpha=0.3); a.legend(fontsize=7)
ab = ax[1,1]
xb = np.arange(2); wbar=0.35
ab.bar(xb-wbar/2, [oldc["meas"], oldc["hid"]], wbar, color="#B23A48", label="old controller")
ab.bar(xb+wbar/2, [v2c["meas"], v2c["hid"]], wbar, color="#1D5FB0", label="controller_v2")
for i,(o,n) in enumerate([(oldc["meas"],v2c["meas"]),(oldc["hid"],v2c["hid"])]):
    ab.text(i-wbar/2, o+0.03, f"{o:.2f}", ha="center", fontsize=8)
    ab.text(i+wbar/2, n+0.03, f"{n:.2f}", ha="center", fontsize=8)
ab.set_xticks(xb); ab.set_xticklabels(["measured","hidden"])
ab.set_title("combined-worst mismatch"); ab.set_ylabel("test RMSE"); ab.grid(axis="y",alpha=0.3); ab.legend(fontsize=8)
fig.suptitle("Model-parameter mismatch robustness: previous controller vs controller_v2\n(frozen nominal-trained observer; v2 degrades far less)", fontsize=13)
plt.tight_layout(rect=[0,0,1,0.96]); plt.savefig("docs/mismatch_v2_compare.png", dpi=130, bbox_inches="tight")
print("saved -> docs/mismatch_v2_compare.png")


# ---- Excel: old vs v2 side by side ----
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Mismatch old vs v2"
WHITE = Font(bold=True, color="FFFFFF"); HEAD = PatternFill("solid", fgColor="1D5FB0")
OLDF = PatternFill("solid", fgColor="F6D3D3"); NEWF = PatternFill("solid", fgColor="D6E6F7")
thin = Side(style="thin", color="C9D2D9"); border = Border(thin,thin,thin,thin)

ws["A1"] = "Model-parameter mismatch — previous controller vs controller_v2"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Frozen nominal-trained observer evaluated on flights with perturbed true parameters (unaware). Test RMSE (avg over 10 unseen flights)."
ws["A2"].font = Font(italic=True, color="555555")

hdr = ["Parameter","Metric","Controller"] + [f"{d:+d}%" if d else "nominal" for d in devs]
for c,h in enumerate(hdr, start=1):
    cell = ws.cell(row=4, column=c, value=h); cell.font = WHITE; cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center"); cell.border = border

r = 5
for p in ["mass","inertia","arm"]:
    for metric,key in [("measured","meas"),("hidden","hid")]:
        for label,src,fillc in [("previous",old,OLDF),("controller_v2",v2,NEWF)]:
            vals = [p, metric, label] + src[p][key]
            for c,val in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=val); cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if c >= 4: cell.number_format = "0.0000"
                if c == 3: cell.fill = fillc
            ws.cell(row=r, column=1).font = Font(bold=True)
            r += 1

ws.cell(row=r+1, column=1, value="Combined-worst (mass -20%, inertia +20%, arm -20%):").font = Font(bold=True)
ws.cell(row=r+2, column=1, value="previous")
ws.cell(row=r+2, column=2, value=f"meas {oldc['meas']:.4f} / hidden {oldc['hid']:.4f}")
ws.cell(row=r+3, column=1, value="controller_v2")
ws.cell(row=r+3, column=2, value=f"meas {v2c['meas']:.4f} / hidden {v2c['hid']:.4f}")
ws.cell(row=r+5, column=1, value="Note: nominal RMSE differs (old 0.0286 vs v2 0.0436), but v2 degrades far less under mismatch — the improved controller keeps flights consistent, so the observer stays robust.").font = Font(italic=True, color="666666")

for col,wd in {"A":15,"B":12,"C":14,"D":10,"E":10,"F":10,"G":10,"H":10}.items(): ws.column_dimensions[col].width = wd
ws.freeze_panes = "D5"
wb.save("docs/mismatch_v2.xlsx")
print("saved -> docs/mismatch_v2.xlsx")
