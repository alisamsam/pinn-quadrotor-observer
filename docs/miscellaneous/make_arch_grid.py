"""Architecture grid: hidden RMSE vs layers (depth) x neurons (width). Heatmap + Excel."""
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

layers=[4,9,12]; neurons=[20,60,100,128]
# hidden RMSE (grid_study.csv); >1 = diverged
H=np.array([[0.1789,0.0631,0.0324,0.0360],
            [0.2536,0.1235,0.0680,0.0740],
            [0.3067,0.1464,2.5233,2.5233]])
Hmeas=np.array([[0.0940,0.0524,0.0266,0.0300],
                [0.1637,0.1207,0.0623,0.0617],
                [0.1946,0.1434,6.8568,5.9885]])
params=np.array([[1792,12552,32912,52876],[3892,30852,83412,135436],[5152,41832,113712,184972]])

# ---- heatmap ----
plot=np.where(H>1,np.nan,H)
fig,ax=plt.subplots(figsize=(8.5,5.5))
im=ax.imshow(plot,cmap="YlOrRd",vmin=0.03,vmax=0.31,aspect="auto")
ax.set_xticks(range(len(neurons))); ax.set_xticklabels(neurons)
ax.set_yticks(range(len(layers))); ax.set_yticklabels(layers)
ax.set_xlabel("neurons per layer (width)"); ax.set_ylabel("hidden layers (depth)")
for i in range(len(layers)):
    for j in range(len(neurons)):
        if H[i,j]>1: ax.text(j,i,f"diverged\n({H[i,j]:.2f})",ha="center",va="center",fontsize=8,color="#333")
        else: ax.text(j,i,f"{H[i,j]:.3f}",ha="center",va="center",fontsize=10,
                      color=("white" if H[i,j]>0.18 else "#222"))
# mark best (4x100)
ax.add_patch(plt.Rectangle((2-0.5,0-0.5),1,1,fill=False,edgecolor="#0B4F8A",lw=3))
ax.text(2,-0.62,"best: 4x100",ha="center",fontsize=9,color="#0B4F8A",fontweight="bold")
fig.colorbar(im,label="hidden-state RMSE")
ax.set_title("Architecture grid — hidden RMSE vs depth x width\n(selection study; lower is better; 4x100 chosen)")
plt.tight_layout(); plt.savefig("docs/arch_grid_heatmap.png",dpi=140,bbox_inches="tight")
print("saved -> docs/arch_grid_heatmap.png")


# ---- Excel ----
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Architecture grid"
WHITE=Font(bold=True,color="FFFFFF"); HEAD=PatternFill("solid",fgColor="1D5FB0")
BEST=PatternFill("solid",fgColor="D6EFD6"); BF=Font(bold=True,color="1B5E20")
BAD=PatternFill("solid",fgColor="F6D3D3")
thin=Side(style="thin",color="C9D2D9"); border=Border(thin,thin,thin,thin); ctr=Alignment(horizontal="center")
ws["A1"]="Architecture grid (selection study): hidden RMSE vs depth x width"; ws["A1"].font=Font(bold=True,size=14)
ws["A2"]="4x100 chosen (lowest hidden RMSE). Deeper nets worse; 12-layer/100-128 diverged. Run on the previous dataset to select the architecture."; ws["A2"].font=Font(italic=True,color="555555")
hdr=["Layers","Neurons","Params","Measured RMSE","Hidden RMSE"]
for c,h in enumerate(hdr,start=1):
    cell=ws.cell(row=4,column=c,value=h); cell.font=WHITE; cell.fill=HEAD; cell.alignment=ctr; cell.border=border
r=5
for i,L in enumerate(layers):
    for j,Nn in enumerate(neurons):
        vals=[L,Nn,int(params[i,j]),round(float(Hmeas[i,j]),4),round(float(H[i,j]),4)]
        for c,v in enumerate(vals,start=1):
            cell=ws.cell(row=r,column=c,value=v); cell.alignment=ctr; cell.border=border
            if c in (4,5): cell.number_format="0.0000"
        if L==4 and Nn==100:
            for c in range(1,6): ws.cell(row=r,column=c).fill=BEST; ws.cell(row=r,column=c).font=BF
        elif H[i,j]>1:
            ws.cell(row=r,column=5).fill=BAD; ws.cell(row=r,column=4).fill=BAD
        r+=1
ws.cell(row=r+1,column=1,value="Note: selection study on the previous dataset; architecture (4x100) reused for the new controller.").font=Font(italic=True,color="666666")
for col,w in {"A":8,"B":9,"C":10,"D":15,"E":14}.items(): ws.column_dimensions[col].width=w
ws.freeze_panes="A5"
wb.save("docs/arch_grid.xlsx"); print("saved -> docs/arch_grid.xlsx")
