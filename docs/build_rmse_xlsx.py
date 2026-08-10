"""Build an Excel of the per-state RMSE results (original vs controller_v2)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# state, unit, group, original RMSE (phase5_results), v2 RMSE (eval_12states_v2)
rows = [
    ("x","m","measured",0.0660,0.0848),
    ("y","m","measured",0.0837,0.0791),
    ("z","m","measured",0.0364,0.0400),
    ("vx","m/s","hidden",0.1880,0.1015),
    ("vy","m/s","hidden",0.2036,0.0926),
    ("vz","m/s","hidden",0.0749,0.0354),
    ("phi","rad","measured",0.0162,0.0232),
    ("theta","rad","measured",0.0274,0.0273),
    ("psi","rad","measured",0.0199,0.0106),
    ("p","rad/s","hidden",0.0217,0.0472),
    ("q","rad/s","hidden",0.0607,0.0516),
    ("r","rad/s","hidden",0.0099,0.0074),
]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "RMSE by state"
NB = Font(bold=True); WHITE = Font(bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1D5FB0")
GOOD = PatternFill("solid", fgColor="D6EFD6"); BAD = PatternFill("solid", fgColor="F6D3D3")
GREY = PatternFill("solid", fgColor="ECEFF1")
thin = Side(style="thin", color="C9D2D9"); border = Border(thin,thin,thin,thin)

ws["A1"] = "PINN Observer — Per-State RMSE (unseen test flights)"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Spiral trajectory R=10, omega=0.6  ·  original controller dataset  vs  improved controller_v2 dataset"
ws["A2"].font = Font(italic=True, color="555555")

hdr = ["State","Unit","Group","Original RMSE","Improved v2 RMSE","Change %","Verdict"]
for c,h in enumerate(hdr, start=1):
    cell = ws.cell(row=4, column=c, value=h); cell.font = WHITE; cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center"); cell.border = border

r0 = 5
for i,(s,u,g,orig,v2) in enumerate(rows):
    r = r0+i
    ch = (v2-orig)/orig*100.0
    verdict = "improved" if v2 < orig else ("worse" if v2 > orig else "same")
    vals = [s,u,g,orig,v2,round(ch,1),verdict]
    for c,val in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=val); cell.border = border
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4).number_format = "0.0000"
    ws.cell(row=r, column=5).number_format = "0.0000"
    ws.cell(row=r, column=6).number_format = "+0.0;-0.0"
    ws.cell(row=r, column=1).font = NB
    ws.cell(row=r, column=6).fill = GOOD if ch < 0 else BAD
    ws.cell(row=r, column=7).fill = GOOD if v2 < orig else BAD


# ---- summary rows ----
summ = [
    ("avg measured (6)", 0.0416, 0.0442),
    ("avg hidden (6)",   0.0931, 0.0559),
    ("avg all (12)",     0.0674, 0.0500),
]
rs = r0 + len(rows) + 1
for j,(name,orig,v2) in enumerate(summ):
    r = rs+j
    ch = (v2-orig)/orig*100.0
    ws.cell(row=r, column=1, value=name).font = NB
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    ws.cell(row=r, column=4, value=orig).number_format = "0.0000"
    ws.cell(row=r, column=5, value=v2).number_format = "0.0000"
    c6 = ws.cell(row=r, column=6, value=round(ch,1)); c6.number_format = "+0.0;-0.0"
    c6.fill = GOOD if ch < 0 else BAD
    for c in range(1,7):
        if c != 6:
            ws.cell(row=r, column=c).fill = GREY
        ws.cell(row=r, column=c).border = border

note = rs + len(summ) + 1
ws.cell(row=note, column=1, value="Notes: RMSE = root-mean-square error between true and estimated state on 10 unseen test flights.")
ws.cell(row=note+1, column=1, value="Negative Change % = improvement (lower error). Same 4x100 observer, weights (0.5,1.5,1.0); only the controller/dataset differs.")
ws.cell(row=note+2, column=1, value="Source: phase5_results.txt (original) and evaluate_12states_v2.py (improved).")
for k in range(3):
    ws.cell(row=note+k, column=1).font = Font(italic=True, color="666666")

widths = {"A":18,"B":8,"C":12,"D":16,"E":18,"F":12,"G":11}
for col,w in widths.items(): ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

out = "docs/observer_v2_rmse.xlsx"
wb.save(out)
print("saved ->", out)
