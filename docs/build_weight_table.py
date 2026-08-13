"""Excel of the loss-weight sweep (7 cases) used to SELECT the observer weights."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# case, w0, w_ode, wy, MSE_0, MSE_g, MSE_y, test_meas_RMSE, test_hidden_RMSE
rows = [
 (1,1.0,1.0,1.0,1.249e-04,6.073e-04,1.389e-03,0.0283,0.0333),
 (2,0.5,1.5,1.0,1.034e-04,4.249e-04,4.644e-04,0.0291,0.0330),
 (3,0.5,0.5,1.0,2.145e-04,1.210e-03,3.124e-03,0.0457,0.0449),
 (4,1.0,2.0,1.0,7.730e-05,5.235e-04,1.507e-03,0.0394,0.0400),
 (5,2.0,1.0,1.0,6.673e-05,7.528e-04,5.515e-04,0.0273,0.0354),
 (6,2.0,1.0,0.5,9.762e-05,6.773e-04,1.989e-03,0.0343,0.0404),
 (7,2.0,1.5,1.5,8.073e-05,7.189e-04,5.870e-04,0.0291,0.0364),
]
hdr = ["Case","w0","w_ode","wy","MSE_0","MSE_g","MSE_y","test meas RMSE","test hidden RMSE"]

wb=openpyxl.Workbook(); ws=wb.active; ws.title="Loss-weight sweep"
WHITE=Font(bold=True,color="FFFFFF"); HEAD=PatternFill("solid",fgColor="1D5FB0")
SEL=PatternFill("solid",fgColor="D6EFD6"); SF=Font(bold=True,color="1B5E20")
thin=Side(style="thin",color="C9D2D9"); border=Border(thin,thin,thin,thin); ctr=Alignment(horizontal="center")

ws["A1"]="Observer loss-weight sweep (4x100) - selection study"; ws["A1"].font=Font(bold=True,size=14)
ws["A2"]="Baseline weights are 1 each. Loss = wy*MSE_y (data) + w_ode*MSE_g (physics) + w0*MSE_0 (initial cond)."; ws["A2"].font=Font(italic=True,color="555555")
ws["A3"]="Selected: Case 2 (0.5, 1.5, 1.0) - lowest hidden RMSE; applied to the improved controller."; ws["A3"].font=Font(italic=True,color="1B5E20")

for c,h in enumerate(hdr,start=1):
    cell=ws.cell(row=5,column=c,value=h); cell.font=WHITE; cell.fill=HEAD; cell.alignment=ctr; cell.border=border
for i,row in enumerate(rows):
    r=6+i
    for c,val in enumerate(row,start=1):
        cell=ws.cell(row=r,column=c,value=val); cell.border=border; cell.alignment=ctr
        if c in (5,6,7): cell.number_format="0.000E+00"
        elif c in (8,9): cell.number_format="0.0000"
    if row[0]==2:
        for c in range(1,10): ws.cell(row=r,column=c).fill=SEL; ws.cell(row=r,column=c).font=SF

ws.cell(row=14,column=1,value="Note: this sweep was run on the previous spiral dataset to choose the weights; the winner (Case 2) was reused for the new controller.").font=Font(italic=True,color="666666")
ws.column_dimensions["A"].width=7
for col in "BCD": ws.column_dimensions[col].width=8
for col in "EFG": ws.column_dimensions[col].width=12
for col in "HI": ws.column_dimensions[col].width=16
ws.freeze_panes="A6"
wb.save("docs/weight_sweep_table.xlsx")
print("saved -> docs/weight_sweep_table.xlsx")
