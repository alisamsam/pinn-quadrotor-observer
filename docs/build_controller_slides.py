"""Controller v2 slide deck: parameters + comparison, and where-used formulas."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DEEP=RGBColor(0x06,0x5A,0x82); TEAL=RGBColor(0x1C,0x72,0x93); MID=RGBColor(0x21,0x29,0x5C)
WHITE=RGBColor(0xFF,0xFF,0xFF); DARK=RGBColor(0x22,0x22,0x22); LTEAL=RGBColor(0xDD,0xEC,0xF2)
GREYT=RGBColor(0x55,0x55,0x55)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def tb(slide,l,t,w,h):
    box=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=box.text_frame
    tf.word_wrap=True; return tf,box
def line(tf,text,size,color,bold=False,font="Calibri",first=False,align=PP_ALIGN.LEFT,italic=False):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.italic=italic; f.name=font; f.color.rgb=color; return p

# ---------- Title slide ----------
s=prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb=MID
tf,_=tb(s,1.0,2.5,11.3,2.6)
line(tf,"Controller v2",54,WHITE,bold=True,font="Cambria",first=True)
line(tf,"Improved robustness for spiral trajectory tracking",26,RGBColor(0xCA,0xDC,0xFC))
tf2,_=tb(s,1.0,6.2,11.3,0.8)
line(tf2,"New tuned gains + saturation limits   |   compared with Controller v1",18,RGBColor(0x9F,0xB6,0xC9),first=True)
line(tf2,"PFE - Laboratoire DRIVE - Competence SIC",14,RGBColor(0x7F,0x93,0xA8))


# ---------- Slide 1: parameters + comparison ----------
s=prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb=WHITE
tf,_=tb(s,0.5,0.35,12.3,0.9)
line(tf,"New Parameters for Robustness (v1 → v2)",34,MID,bold=True,font="Cambria",first=True)

# comparison table (left)
rows=[("Gain","v1","v2"),("KP_XY","0.6","4.5"),("KD_XY","1.0","5.0"),("KP_Z","4.0","6.0"),
      ("KD_Z","3.0","5.0"),("KP_ATT","8.0","25"),("KD_ATT","4.0","8"),("KP_yaw","2.0","2.0"),("KD_yaw","2.0","2.0")]
tbl=s.shapes.add_table(len(rows),3,Inches(0.6),Inches(1.5),Inches(5.6),Inches(4.9)).table
tbl.columns[0].width=Inches(2.6); tbl.columns[1].width=Inches(1.5); tbl.columns[2].width=Inches(1.5)
for ci in range(3):
    c=tbl.cell(0,ci); c.fill.solid(); c.fill.fore_color.rgb=TEAL
    r=c.text_frame.paragraphs[0]; run=r.add_run(); run.text=rows[0][ci]; run.font.bold=True; run.font.color.rgb=WHITE; run.font.size=Pt(16); run.font.name="Calibri"; r.alignment=PP_ALIGN.CENTER
for ri in range(1,len(rows)):
    for ci in range(3):
        c=tbl.cell(ri,ci); c.fill.solid(); c.fill.fore_color.rgb=(LTEAL if ci==2 else WHITE)
        r=c.text_frame.paragraphs[0]; run=r.add_run(); run.text=rows[ri][ci]
        run.font.size=Pt(15); run.font.name="Calibri"; run.font.color.rgb=DARK
        run.font.bold=(ci==2); r.alignment=(PP_ALIGN.LEFT if ci==0 else PP_ALIGN.CENTER)

# saturation box (right)
box=s.shapes.add_shape(1,Inches(6.7),Inches(1.5),Inches(6.1),Inches(2.9))
box.fill.solid(); box.fill.fore_color.rgb=RGBColor(0xF0,0xF5,0xF8); box.line.color.rgb=TEAL; box.line.width=Pt(1.5)
tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.18)
line(tf,"Safety limits (saturation) — new in v2",18,DEEP,bold=True,font="Cambria",first=True)
for t in ["Thrust  U1 ∈ [1, 35] N   (≤ 2× weight; no downward pull)",
          "Tilt  |φ_d|, |θ_d|  ≤ 30°",
          "Torques  U2, U3, U4 ∈ [-5, +5]"]:
    p=line(tf,t,15,DARK); p.level=0
tf2,_=tb(s,6.7,4.6,6.1,1.7)
line(tf2,"Also new (method upgrades):",16,DEEP,bold=True,font="Cambria",first=True)
line(tf2,"• acceleration feedforward   • thrust tilt-compensation   • exact tilt inversion",15,DARK)
line(tf2,"Result: steady-state tracking ~5 m (v1) → ~0.18 m (v2), within limits.",15,TEAL,italic=True)


# ---------- Slide 2: where the constants are used (formulas) ----------
s=prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb=WHITE
tf,_=tb(s,0.5,0.35,12.3,0.9)
line(tf,"Where the Constants Are Used — Formulas",34,MID,bold=True,font="Cambria",first=True)

steps=[
 ("1 · Horizontal loop  (feedforward + PD)",
  "a_x = a_d,x + KP_XY·(x_d - x) + KD_XY·(v_d,x - v_x)     (a_y analogous)",
  "uses  KP_XY = 4.5,  KD_XY = 5.0"),
 ("2 · Thrust with tilt-compensation   (saturated to [1, 35] N)",
  "U1 = m·( g + a_d,z + KP_Z·(z_d - z) + KD_Z·(v_d,z - v_z) ) / (cosφ · cosθ)",
  "uses  KP_Z = 6.0,  KD_Z = 5.0"),
 ("3 · Desired tilt  (exact inversion, saturated to ≤ 30°)",
  "φ_d = arcsin( U_ex·sinψ - U_ey·cosψ ),    θ_d = arcsin( ... / cosφ_d )",
  "uses  tilt limit 30°"),
 ("4 · Attitude torques   (saturated to [-5, +5])",
  "U2 = (I_x / l)·( KP_ATT·(φ_d - φ) - KD_ATT·p - coupling )    (U3, U4 analogous)",
  "uses  KP_ATT = 25,  KD_ATT = 8,  KP_yaw = 2,  KD_yaw = 2"),
]
y=1.45
for head,formula,uses in steps:
    tf,_=tb(s,0.6,y,12.2,1.25)
    line(tf,head,17,DEEP,bold=True,font="Cambria",first=True)
    line(tf,formula,16,DARK,font="Cambria")
    line(tf,uses,14,TEAL,bold=True)
    y+=1.36

prs.save("docs/Controller_v2_slides.pptx")
print("saved -> docs/Controller_v2_slides.pptx")
