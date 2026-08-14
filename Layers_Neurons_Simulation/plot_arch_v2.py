"""Architecture-grid histogram + Excel for the NEW controller (v2).
Grouped bars by depth, log-scale y (same style as the previous-controller grid).
Outputs into this folder (Layers_Neurons_Simulation/).
"""
import os, csv
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
CSV  = os.path.join(HERE, "grid_study_v2.csv")
DIVERGE = 1.0   # rmse above this = training diverged

rows = list(csv.DictReader(open(CSV)))
data = {}
for r in rows:
    L, H = int(r["layers"]), int(r["neurons"])
    data.setdefault(L, {})[H] = (float(r["rmse_meas"]), float(r["rmse_hidden"]), int(r["params"]))
LAYERS  = sorted(data)
NEURONS = sorted({int(r["neurons"]) for r in rows})
COLORS  = {4: "#1565C0", 9: "#2E9E4F", 12: "#D62728"}   # blue / green / red

# best architecture by hidden RMSE among converged runs
conv = [(int(r["layers"]), int(r["neurons"]), float(r["rmse_hidden"]))
        for r in rows if float(r["rmse_hidden"]) < DIVERGE]
bestL, bestH, best_val = min(conv, key=lambda t: t[2])

# ---------------- histogram ----------------
x = np.arange(len(NEURONS)); n = len(LAYERS); w = 0.8 / n
fig, ax = plt.subplots(figsize=(11.5, 6.4))
for i, L in enumerate(LAYERS):
    vals = [data[L].get(H, (np.nan, np.nan, 0))[1] for H in NEURONS]
    off  = (i - (n - 1) / 2) * w
    bars = ax.bar(x + off, vals, w, label=f"{L} layers",
                  color=COLORS.get(L, "#888888"), edgecolor="white", linewidth=0.6)
    for b, v in zip(bars, vals):
        if np.isfinite(v):
            txt = f"{v:.3f}"
            ax.annotate(txt, (b.get_x() + b.get_width()/2, v), xytext=(0, 3),
                        textcoords="offset points", ha="center", va="bottom",
                        fontsize=8, rotation=90)
ax.set_yscale("log")
ax.set_xticks(x); ax.set_xticklabels(NEURONS)
ax.set_xlabel("Neurons per layer", fontsize=11)
ax.set_ylabel("Hidden-state RMSE  (log scale)", fontsize=11)
ax.set_title("Architecture grid: hidden RMSE vs width and depth (neutral weights)",
             fontsize=12)
ax.grid(axis="y", which="both", ls="-", lw=0.4, alpha=0.35); ax.set_axisbelow(True)
ax.legend(title="Depth", loc="upper left", framealpha=0.95); ax.margins(y=0.22)
fig.tight_layout()
PNG = os.path.join(HERE, "arch_grid_v2.png")
fig.savefig(PNG, dpi=150, bbox_inches="tight"); plt.close(fig)
print("saved", PNG)

# ---------------- Excel ----------------
wb = Workbook(); ws = wb.active; ws.title = "arch_grid_v2"
hdr = ["Layers","Neurons","Params","Measured RMSE","Hidden RMSE","Train time (s)","Note"]
navy = PatternFill("solid", fgColor="1F2A5A"); green = PatternFill("solid", fgColor="4E7D3A")
red  = PatternFill("solid", fgColor="C0392B"); thin = Side(style="thin", color="BBBBBB")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.append(hdr)
for c in range(1, len(hdr)+1):
    cell = ws.cell(1, c); cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = navy; cell.alignment = Alignment(horizontal="center"); cell.border = bd
for r in rows:
    L, H = int(r["layers"]), int(r["neurons"])
    m, h, p = float(r["rmse_meas"]), float(r["rmse_hidden"]), int(r["params"])
    diverged = (m >= DIVERGE or h >= DIVERGE)
    best = (L == bestL and H == bestH)
    note = "diverged" if diverged else ("BEST (lowest hidden)" if best else "")
    ws.append([L, H, p, m, h, int(r["train_time_s"]), note]); row = ws.max_row
    for c in range(1, len(hdr)+1):
        cell = ws.cell(row, c); cell.alignment = Alignment(horizontal="center"); cell.border = bd
    if diverged:
        for c in (4, 5): ws.cell(row, c).fill = red; ws.cell(row, c).font = Font(bold=True, color="FFFFFF")
    if best:
        ws.cell(row, 5).fill = green; ws.cell(row, 5).font = Font(bold=True, color="FFFFFF")
widths = [9, 9, 10, 15, 13, 15, 22]
for i, wd in enumerate(widths, start=1):
    ws.column_dimensions[chr(64+i)].width = wd
XLSX = os.path.join(HERE, "arch_grid_v2_table.xlsx")
wb.save(XLSX); print("saved", XLSX)
print(f"BEST by hidden RMSE: {bestL}x{bestH} = {best_val:.4f}")
