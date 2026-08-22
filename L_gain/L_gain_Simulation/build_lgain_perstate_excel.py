"""Excel table for the per-state L-OFF vs L-ON comparison (reads lgain_perstate_v2.csv)."""
import os, csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
rows = list(csv.DictReader(open(os.path.join(HERE, "lgain_perstate_v2.csv"))))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "L-gain per state"
WHITE = Font(bold=True, color="FFFFFF"); HEAD = PatternFill("solid", fgColor="1F2A5A")
hidfill = PatternFill("solid", fgColor="E8EEF7"); bad = PatternFill("solid", fgColor="C0392B")
thin = Side(style="thin", color="C9D2D9"); bd = Border(thin,thin,thin,thin)

ws["A1"] = "Per-state observer RMSE - without vs with the learned gain L (unseen test flights)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("without L = plain PINN observer (4x100); with L = PINN + learned gain L*(y - yhat). "
            "Ratio = with L / without L; >1 means L made it worse.")
ws["A2"].font = Font(italic=True, color="555555")

hdr = ["State","Unit","Type","RMSE without L","RMSE with L","Ratio (with / without)"]
for c,h in enumerate(hdr, start=1):
    cell = ws.cell(row=4, column=c, value=h); cell.font = WHITE; cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center"); cell.border = bd

moff=mon=hoff=hon=0.0; nm=nh=0
r = 5
for d in rows:
    off, on = float(d["rmse_without_L"]), float(d["rmse_with_L"])
    ratio = on/off if off else 0.0
    hidden = d["type"] == "hidden"
    for c,val in enumerate([d["state"], d["unit"], d["type"], off, on, round(ratio,2)], start=1):
        cell = ws.cell(row=r, column=c, value=val); cell.border = bd
        cell.alignment = Alignment(horizontal="center")
        if hidden: cell.fill = hidfill
    ws.cell(row=r, column=4).number_format = "0.0000"
    ws.cell(row=r, column=5).number_format = "0.0000"
    ws.cell(row=r, column=6).number_format = "0.0"
    if ratio >= 3:   # big blow-up -> red ratio
        ws.cell(row=r, column=6).fill = bad; ws.cell(row=r, column=6).font = Font(bold=True, color="FFFFFF")
    if hidden: hoff+=off; hon+=on; nh+=1
    else:      moff+=off; mon+=on; nm+=1
    r += 1

# average rows
def avg_row(label, ao, an, rr):
    global r
    for c,val in enumerate([label,"","", round(ao,4), round(an,4), round(an/ao,2)], start=1):
        cell = ws.cell(row=rr, column=c, value=val); cell.border = bd
        cell.alignment = Alignment(horizontal="center"); cell.font = Font(bold=True)
    ws.cell(row=rr, column=4).number_format = "0.0000"
    ws.cell(row=rr, column=5).number_format = "0.0000"
    ws.cell(row=rr, column=6).number_format = "0.0"
avg_row("avg measured", moff/nm, mon/nm, r); r+=1
avg_row("avg hidden",   hoff/nh, hon/nh, r); r+=1

ws.cell(row=r+1, column=1, value=("Finding: L barely changes the measured states but inflates the hidden states ~6x "
                                  "(0.056 -> 0.343); the fast rates p, q and velocities vx, vy are hit hardest. "
                                  "The plain PINN observer (without L) is retained."))
ws.cell(row=r+1, column=1).font = Font(italic=True, color="666666")
for col,wd in {"A":16,"B":8,"C":11,"D":16,"E":14,"F":20}.items(): ws.column_dimensions[col].width = wd
ws.freeze_panes = "A5"
XLSX = os.path.join(HERE, "lgain_perstate_v2_table.xlsx")
wb.save(XLSX); print("saved", XLSX)
