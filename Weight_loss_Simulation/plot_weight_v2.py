"""Loss-weight sensitivity histogram for the NEW controller (v2), 4x100.
Same style as the previous-controller figure: grouped bars (measured/hidden),
noise band above best hidden, rotated value labels.
Outputs into this folder (Weight_loss_Simulation/).
"""
import os, csv
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE  = os.path.dirname(os.path.abspath(__file__))
CSV   = os.path.join(HERE, "weight_study_v2.csv")
NOISE = 0.009   # run-to-run std from the seed study (same as v1)

rows = list(csv.DictReader(open(CSV)))
cases    = [f"C{r['case']}\n({r['w0']},{r['w_ode']},{r['wy']})" for r in rows]
measured = [float(r["test_meas_RMSE"])   for r in rows]
hidden   = [float(r["test_hidden_RMSE"]) for r in rows]
best_h   = min(hidden)

# ---------------- histogram (exact v1 style) ----------------
x = np.arange(len(rows)); w = 0.38
fig, ax = plt.subplots(figsize=(11, 6))
b1 = ax.bar(x - w/2, measured, w, label="measured RMSE", color="#E08E0B")
b2 = ax.bar(x + w/2, hidden,   w, label="hidden RMSE",   color="#0B3C5D")
ax.axhspan(best_h, best_h + NOISE, color="grey", alpha=0.22,
           label=f"±{NOISE:.3f} noise band above best")
for b, v in zip(b1, measured):
    ax.text(b.get_x()+b.get_width()/2, v+0.001, f"{v:.4f}",
            ha="center", va="bottom", fontsize=8, rotation=90, color="#8A5300")
for b, v in zip(b2, hidden):
    ax.text(b.get_x()+b.get_width()/2, v+0.001, f"{v:.4f}",
            ha="center", va="bottom", fontsize=8, rotation=90, color="#0B3C5D")
ax.set_xticks(x); ax.set_xticklabels(cases, fontsize=8)
ax.set_xlabel("Weight case  (w0, w_ode, wy)")
ax.set_ylabel("RMSE on unseen flights")
ax.set_title("Loss-weight sensitivity at 4x100 (new controller v2)")
ax.legend(); ax.grid(axis="y", alpha=0.3)
ax.set_ylim(0, max(max(hidden), max(measured)) * 1.25)
fig.tight_layout()
PNG = os.path.join(HERE, "weight_sensitivity_v2.png")
fig.savefig(PNG, dpi=140, bbox_inches="tight"); plt.close(fig)
print("saved", PNG)

# ---------------- Excel table ----------------
wb = Workbook(); ws = wb.active; ws.title = "weight_study_v2"
hdr = ["Case","w0","w_ode","wy","MSE_0","MSE_g","MSE_y","Measured RMSE","Hidden RMSE","In noise band?"]
navy = PatternFill("solid", fgColor="1F2A5A"); green = PatternFill("solid", fgColor="4E7D3A")
band = PatternFill("solid", fgColor="E8EEF7"); thin = Side(style="thin", color="BBBBBB")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.append(hdr)
for c in range(1, len(hdr)+1):
    cell = ws.cell(1, c); cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = navy; cell.alignment = Alignment(horizontal="center"); cell.border = bd
best_meas = min(measured); last = 1
for i, r in enumerate(rows):
    m, h = measured[i], hidden[i]
    inband = h <= best_h + NOISE
    ws.append([f"C{r['case']}", float(r['w0']), float(r['w_ode']), float(r['wy']),
               float(r['MSE_0']), float(r['MSE_g']), float(r['MSE_y']),
               m, h, "yes" if inband else ""])
    last = ws.max_row
    for c in range(1, len(hdr)+1):
        cell = ws.cell(last, c); cell.alignment = Alignment(horizontal="center"); cell.border = bd
        if inband: cell.fill = band
    if h == best_h:
        ws.cell(last, 9).fill = green; ws.cell(last, 9).font = Font(bold=True, color="FFFFFF")
    if m == best_meas:
        ws.cell(last, 8).fill = green; ws.cell(last, 8).font = Font(bold=True, color="FFFFFF")
ws.cell(last+2, 1).value = ("Selected: C2 (0.5, 1.5, 1.0) - best measured RMSE and inside the "
                            "noise band on hidden; C4 has the single lowest hidden but is "
                            "statistically tied with C2, C3, C7.")
ws.cell(last+2, 1).font = Font(italic=True)
widths = [16,7,8,7,11,11,11,15,13,15]
for i, wdt in enumerate(widths, start=1):
    ws.column_dimensions[chr(64+i)].width = wdt
XLSX = os.path.join(HERE, "weight_study_v2_table.xlsx")
wb.save(XLSX); print("saved", XLSX)
