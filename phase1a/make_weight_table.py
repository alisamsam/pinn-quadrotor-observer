import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

data = [
 [1,1.0,1.0,1.0,"5.71e-04","8.49e-04","4.52e-03",0.0509,0.0411],
 [2,0.5,1.5,1.0,"7.75e-05","3.45e-04","3.27e-04",0.0239,0.0275],
 [3,0.5,0.5,1.0,"1.20e-04","1.00e-03","1.73e-03",0.0416,0.0368],
 [4,1.0,2.0,1.0,"8.36e-05","3.06e-04","3.02e-04",0.0239,0.0272],
 [5,2.0,1.0,1.0,"1.61e-04","1.14e-03","2.75e-03",0.0547,0.0455],
 [6,2.0,1.0,0.5,"4.83e-05","6.02e-04","7.08e-04",0.0337,0.0354],
 [7,2.0,1.5,1.5,"1.95e-04","8.09e-04","1.28e-03",0.0279,0.0356],
]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Weight study"
NAVY="21295C"; GREEN="D4EDDA"
ws["A1"]="Loss-weight sensitivity study (no-gain [t,x0] observer) - test on unseen flights"
ws["A1"].font=Font(name="Arial",bold=True,size=12); ws.merge_cells("A1:I1")
hdr=["Case","w0","w_ode","w_y","MSE_0","MSE_g","MSE_y","meas RMSE","hidden RMSE"]
for j,h in enumerate(hdr,1):
    c=ws.cell(row=3,column=j,value=h)
    c.font=Font(name="Arial",bold=True,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=NAVY); c.alignment=Alignment(horizontal="center",wrap_text=True)
thin=Side(style="thin",color="CCCCCC"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
best=min(r[8] for r in data)
for i,r in enumerate(data,start=4):
    for j,v in enumerate(r,1):
        c=ws.cell(row=i,column=j,value=v)
        c.font=Font(name="Arial"); c.border=border; c.alignment=Alignment(horizontal="center")
        if j==9 and abs(r[8]-best)<1e-9: c.fill=PatternFill("solid",fgColor=GREEN)
for j,w in enumerate([6,6,7,6,12,12,12,11,12],1): ws.column_dimensions[chr(64+j)].width=w
ws.row_dimensions[3].height=30; ws.freeze_panes="A4"
wb.save("docs/weight_study.xlsx")
print("wrote docs/weight_study.xlsx")